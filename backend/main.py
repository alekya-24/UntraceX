import os
import shutil
import uuid
import json
from datetime import datetime, timezone
from typing import List, Optional
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form, status, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from database import engine, Base, get_db, SessionLocal, UPLOADS_DIR, CLEANED_DIR
from models import User, Document, DocumentMetadata, DocumentAnalysis, CleaningAction
from auth import (
    hash_password,
    verify_password,
    create_access_token,
    get_current_user
)
from document_cleaner import DocumentProcessor
from ai_detector import analyze_linguistic_patterns
from text_improver import improve_text, generate_diff_chunks

import pymupdf
import docx
import pptx

# Ensure database tables exist
Base.metadata.create_all(bind=engine)

app = FastAPI(title="UntraceX API", version="2.0.0")

# Enable CORS for frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
def startup_seed():
    db = SessionLocal()
    try:
        demo_user = db.query(User).filter(User.email == "demo@untracex.local").first()
        if not demo_user:
            demo = User(
                name="Demo User",
                email="demo@untracex.local",
                password_hash=hash_password("demo1234")
            )
            db.add(demo)
            db.commit()
    finally:
        db.close()

# ----------------- SCHEMAS -----------------
class RegisterRequest(BaseModel):
    name: str
    email: str
    password: str

class LoginRequest(BaseModel):
    email: str
    password: str

class CleanRequest(BaseModel):
    keys_to_remove: List[str]

class ImproveRequest(BaseModel):
    custom_text: Optional[str] = None
    target_sentences: Optional[List[str]] = None

# ----------------- AUTH ENDPOINTS -----------------
@app.post("/api/auth/register", status_code=status.HTTP_201_CREATED)
def register(req: RegisterRequest, db: Session = Depends(get_db)):
    existing = db.query(User).filter(User.email == req.email.lower().strip()).first()
    if existing:
        raise HTTPException(status_code=400, detail="Email is already registered.")
    
    if len(req.password) < 6:
        raise HTTPException(status_code=400, detail="Password must be at least 6 characters.")

    new_user = User(
        name=req.name.strip(),
        email=req.email.lower().strip(),
        password_hash=hash_password(req.password)
    )
    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    token = create_access_token(data={"sub": new_user.id, "email": new_user.email})
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": new_user.id,
            "name": new_user.name,
            "email": new_user.email,
            "created_at": new_user.created_at.isoformat()
        }
    }

@app.post("/api/auth/login")
def login(req: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == req.email.lower().strip()).first()
    if not user or not verify_password(req.password, user.password_hash):
        raise HTTPException(status_code=400, detail="Invalid email or password.")

    token = create_access_token(data={"sub": user.id, "email": user.email})
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {
            "id": user.id,
            "name": user.name,
            "email": user.email,
            "created_at": user.created_at.isoformat()
        }
    }

@app.get("/api/auth/me")
def get_me(current_user: User = Depends(get_current_user)):
    return {
        "id": current_user.id,
        "name": current_user.name,
        "email": current_user.email,
        "created_at": current_user.created_at.isoformat()
    }

# ----------------- DOCUMENT ENDPOINTS -----------------

@app.post("/api/documents/upload", status_code=status.HTTP_201_CREATED)
async def upload_document(
    file: Optional[UploadFile] = File(None),
    pasted_text: Optional[str] = Form(None),
    filename_override: Optional[str] = Form(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc_id = str(uuid.uuid4())

    if file:
        filename = file.filename or "uploaded_document"
        ext = filename.split(".")[-1].upper() if "." in filename else "TXT"
        if ext not in ["PDF", "DOCX", "PPTX", "TXT"]:
            raise HTTPException(
                status_code=400, 
                detail=f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, PPTX, TXT."
            )
        
        saved_filename = f"{doc_id}_{filename}"
        file_path = os.path.join(UPLOADS_DIR, saved_filename)
        
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        file_size = os.path.getsize(file_path)
        file_type = ext

    elif pasted_text:
        title = filename_override.strip() if filename_override else "pasted_document.txt"
        if not title.lower().endswith(".txt"):
            title += ".txt"
        filename = title
        file_type = "TXT"
        saved_filename = f"{doc_id}_{filename}"
        file_path = os.path.join(UPLOADS_DIR, saved_filename)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(pasted_text)
        file_size = os.path.getsize(file_path)

    else:
        raise HTTPException(status_code=400, detail="Please provide a document file or paste text content.")

    new_doc = Document(
        id=doc_id,
        user_id=current_user.id,
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        file_path=file_path,
        status="Uploaded"
    )
    db.add(new_doc)
    db.commit()
    db.refresh(new_doc)

    return {
        "id": new_doc.id,
        "filename": new_doc.filename,
        "file_type": new_doc.file_type,
        "file_size": new_doc.file_size,
        "status": new_doc.status,
        "created_at": new_doc.created_at.isoformat()
    }


@app.get("/api/documents")
def list_documents(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    docs = db.query(Document).filter(Document.user_id == current_user.id).order_by(Document.created_at.desc()).all()
    result = []
    for d in docs:
        latest_analysis = db.query(DocumentAnalysis).filter(DocumentAnalysis.document_id == d.id).order_by(DocumentAnalysis.created_at.desc()).first()
        result.append({
            "id": d.id,
            "filename": d.filename,
            "file_type": d.file_type,
            "file_size": d.file_size,
            "status": d.status,
            "ai_likelihood": latest_analysis.ai_likelihood if latest_analysis else None,
            "indicator_level": latest_analysis.indicator_level if latest_analysis else None,
            "has_cleaned": bool(d.cleaned_file_path and os.path.exists(d.cleaned_file_path)),
            "created_at": d.created_at.isoformat()
        })
    return result


@app.get("/api/documents/{id}")
def get_document(id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    metadata_records = db.query(DocumentMetadata).filter(DocumentMetadata.document_id == doc.id).all()
    latest_analysis = db.query(DocumentAnalysis).filter(DocumentAnalysis.document_id == doc.id).order_by(DocumentAnalysis.created_at.desc()).first()
    cleaning_action = db.query(CleaningAction).filter(CleaningAction.document_id == doc.id).order_by(CleaningAction.created_at.desc()).first()

    analysis_data = None
    if latest_analysis:
        analysis_data = {
            "ai_likelihood": latest_analysis.ai_likelihood,
            "human_likelihood": 100 - latest_analysis.ai_likelihood,
            "level": latest_analysis.indicator_level,
            "signals": json.loads(latest_analysis.signals) if latest_analysis.signals else [],
            "explanation": latest_analysis.explanation
        }

    cleaning_summary = None
    if cleaning_action:
        try:
            cleaning_summary = {
                "removed_metadata": json.loads(cleaning_action.removed_metadata) if cleaning_action.removed_metadata.startswith("[") else cleaning_action.removed_metadata.split(","),
                "before": json.loads(cleaning_action.before_state) if cleaning_action.before_state else {},
                "after": json.loads(cleaning_action.after_state) if cleaning_action.after_state else {},
                "text_changes": json.loads(cleaning_action.text_changes) if cleaning_action.text_changes else None,
                "created_at": cleaning_action.created_at.isoformat()
            }
        except Exception:
            cleaning_summary = {
                "removed_metadata": cleaning_action.removed_metadata,
                "created_at": cleaning_action.created_at.isoformat()
            }

    return {
        "id": doc.id,
        "filename": doc.filename,
        "file_type": doc.file_type,
        "file_size": doc.file_size,
        "status": doc.status,
        "extracted_text": doc.extracted_text,
        "cleaned_text": doc.cleaned_text,
        "has_cleaned": bool(doc.cleaned_file_path and os.path.exists(doc.cleaned_file_path)),
        "created_at": doc.created_at.isoformat(),
        "metadata": [
            {
                "id": m.id,
                "name": m.field_name,
                "value": m.field_value,
                "removable": m.removable,
                "removed": m.removed
            }
            for m in metadata_records
        ],
        "analysis": analysis_data,
        "cleaning_summary": cleaning_summary
    }


@app.post("/api/documents/{id}/analyze")
def analyze_document_endpoint(id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    if not os.path.exists(doc.file_path):
        raise HTTPException(status_code=404, detail="Document source file missing on server.")

    # 1. Extract metadata and visible text
    try:
        metadata_list, extracted_text = DocumentProcessor.inspect(doc.file_path, doc.file_type)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to extract document content: {str(e)}")

    doc.extracted_text = extracted_text

    # 2. Save metadata entries
    db.query(DocumentMetadata).filter(DocumentMetadata.document_id == doc.id).delete()
    for item in metadata_list:
        meta_entry = DocumentMetadata(
            document_id=doc.id,
            metadata_key=item["key"],
            field_name=item["key"],
            metadata_value=item["value"],
            field_value=item["value"],
            removable=item["removable"],
            removed=False
        )
        db.add(meta_entry)

    # 3. Analyze text for AI-related writing indicators
    nlp_results = analyze_linguistic_patterns(extracted_text)

    # 4. Save analysis to database
    db.query(DocumentAnalysis).filter(DocumentAnalysis.document_id == doc.id).delete()
    analysis_record = DocumentAnalysis(
        document_id=doc.id,
        ai_likelihood=nlp_results["ai_likelihood"],
        indicator_level=nlp_results["level"],
        signals=json.dumps(nlp_results["signals"]),
        explanation=nlp_results["explanation"]
    )
    db.add(analysis_record)

    doc.status = "Scanned" if doc.status != "Cleaned" else "Cleaned"
    db.commit()

    # 5. Return structured response matching Section 22 specification
    return {
        "ai_likelihood": nlp_results["ai_likelihood"],
        "human_likelihood": nlp_results["human_likelihood"],
        "level": nlp_results["level"],
        "level_label": nlp_results["level_label"],
        "signals": nlp_results["signals"],
        "metadata": [
            {
                "name": item["key"],
                "value": item["value"],
                "removable": item["removable"]
            }
            for item in metadata_list
        ],
        "extracted_text": extracted_text,
        "explanation": nlp_results["explanation"]
    }


@app.get("/api/documents/{id}/analysis")
def get_analysis_endpoint(id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    latest_analysis = db.query(DocumentAnalysis).filter(DocumentAnalysis.document_id == doc.id).order_by(DocumentAnalysis.created_at.desc()).first()
    if not latest_analysis:
        return analyze_document_endpoint(id, current_user, db)

    metadata_records = db.query(DocumentMetadata).filter(DocumentMetadata.document_id == doc.id).all()
    signals = json.loads(latest_analysis.signals) if latest_analysis.signals else []

    return {
        "ai_likelihood": latest_analysis.ai_likelihood,
        "human_likelihood": 100 - latest_analysis.ai_likelihood,
        "level": latest_analysis.indicator_level,
        "signals": signals,
        "metadata": [
            {
                "name": m.field_name,
                "value": m.field_value,
                "removable": m.removable
            }
            for m in metadata_records
        ],
        "extracted_text": doc.extracted_text,
        "explanation": latest_analysis.explanation
    }


# Backwards compatibility endpoint alias for scan
@app.post("/api/documents/{id}/scan")
def scan_document(id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return analyze_document_endpoint(id, current_user, db)


@app.post("/api/documents/{id}/clean")
def clean_document_endpoint(
    id: str,
    req: CleanRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    if not os.path.exists(doc.file_path):
        raise HTTPException(status_code=404, detail="Source file missing on server.")

    if not req.keys_to_remove:
        raise HTTPException(status_code=400, detail="No metadata keys selected for removal.")

    cleaned_filename = f"cleaned_{doc.id}_{doc.filename}"
    cleaned_path = os.path.join(CLEANED_DIR, cleaned_filename)

    try:
        summary = DocumentProcessor.clean(
            file_path=doc.file_path,
            file_type=doc.file_type,
            cleaned_path=cleaned_path,
            keys_to_remove=req.keys_to_remove
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to clean document: {str(e)}")

    doc.cleaned_file_path = cleaned_path
    doc.status = "Cleaned"

    # Mark metadata items as removed
    for m in db.query(DocumentMetadata).filter(DocumentMetadata.document_id == doc.id).all():
        if m.field_name in req.keys_to_remove:
            m.removed = True

    # Record cleaning action
    action = CleaningAction(
        document_id=doc.id,
        removed_fields=json.dumps(req.keys_to_remove),
        removed_metadata=json.dumps(req.keys_to_remove),
        before_state=json.dumps(summary["before"]),
        after_state=json.dumps(summary["after"]),
        text_changes=None
    )
    db.add(action)
    db.commit()

    return {
        "message": "Supported metadata removed successfully.",
        "status": "Cleaned",
        "cleaned_filename": doc.filename,
        "summary": summary
    }


@app.post("/api/documents/{id}/improve")
def improve_text_endpoint(
    id: str,
    req: ImproveRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    text_to_process = req.custom_text if req.custom_text is not None else (doc.extracted_text or "")
    if not text_to_process.strip():
        raise HTTPException(status_code=400, detail="No text content to improve.")

    # Apply text improvement algorithm
    improved_result = improve_text(text_to_process, req.target_sentences)
    new_text = improved_result["improved_text"]
    doc.cleaned_text = new_text

    # Re-evaluate AI likelihood on improved text
    new_analysis = analyze_linguistic_patterns(new_text)

    # Save to cleaned file path
    cleaned_filename = f"cleaned_{doc.id}_{doc.filename}"
    cleaned_path = os.path.join(CLEANED_DIR, cleaned_filename)

    if doc.file_type.upper() == "TXT":
        with open(cleaned_path, "w", encoding="utf-8") as f:
            f.write(new_text)
        doc.cleaned_file_path = cleaned_path
    elif doc.file_type.upper() == "DOCX" and os.path.exists(doc.file_path):
        try:
            d = docx.Document(doc.file_path)
            # update paragraphs with improved text
            paragraphs = new_text.splitlines()
            for idx, p_text in enumerate(paragraphs):
                if idx < len(d.paragraphs):
                    d.paragraphs[idx].text = p_text
                else:
                    d.add_paragraph(p_text)
            d.save(cleaned_path)
            doc.cleaned_file_path = cleaned_path
        except Exception:
            pass
    elif doc.file_type.upper() == "PDF" and os.path.exists(doc.file_path):
        try:
            # Create a clean sanitized PDF rendering the improved text
            new_doc = pymupdf.open()
            page = new_doc.new_page()
            rect = pymupdf.Rect(50, 50, 550, 750)
            page.insert_textbox(rect, new_text, fontsize=10, fontname="helv")
            new_doc.save(cleaned_path, clean=True, deflate=True)
            new_doc.close()
            doc.cleaned_file_path = cleaned_path
        except Exception:
            pass

    doc.status = "Cleaned"

    # Save cleaning action with text changes
    action = CleaningAction(
        document_id=doc.id,
        removed_fields=json.dumps(["AI-related writing pattern refinement"]),
        removed_metadata=json.dumps(["AI-related writing pattern refinement"]),
        text_changes=json.dumps(improved_result["diff_chunks"]),
        before_state=json.dumps({"text_sample": text_to_process[:200]}),
        after_state=json.dumps({"text_sample": new_text[:200]})
    )
    db.add(action)
    db.commit()

    return {
        "message": "Writing pattern improved successfully.",
        "original_text": text_to_process,
        "improved_text": new_text,
        "diff_chunks": improved_result["diff_chunks"],
        "new_analysis": new_analysis
    }


@app.get("/api/documents/{id}/compare")
def compare_document_endpoint(id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    original_text = doc.extracted_text or ""
    cleaned_text = doc.cleaned_text or original_text

    diff_chunks = generate_diff_chunks(original_text, cleaned_text)

    # Get metadata before vs after
    metadata_records = db.query(DocumentMetadata).filter(DocumentMetadata.document_id == doc.id).all()
    meta_before = {m.field_name: m.field_value for m in metadata_records}
    meta_after = {m.field_name: ("Removed" if m.removed else m.field_value) for m in metadata_records}

    return {
        "original_text": original_text,
        "cleaned_text": cleaned_text,
        "diff_chunks": diff_chunks,
        "metadata_before": meta_before,
        "metadata_after": meta_after,
        "has_cleaned": bool(doc.cleaned_file_path and os.path.exists(doc.cleaned_file_path))
    }


@app.get("/api/documents/{id}/download")
def download_document(
    id: str,
    version: str = Query("cleaned", description="'cleaned' or 'original'"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    if version == "cleaned":
        if not doc.cleaned_file_path or not os.path.exists(doc.cleaned_file_path):
            raise HTTPException(status_code=400, detail="Cleaned version has not been created yet.")
        file_to_send = doc.cleaned_file_path
        dl_filename = f"untraced_{doc.filename}"
    else:
        if not os.path.exists(doc.file_path):
            raise HTTPException(status_code=404, detail="Original file missing.")
        file_to_send = doc.file_path
        dl_filename = doc.filename

    return FileResponse(
        path=file_to_send,
        filename=dl_filename,
        media_type="application/octet-stream"
    )


@app.delete("/api/documents/{id}")
def delete_document(id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = db.query(Document).filter(Document.id == id, Document.user_id == current_user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    if os.path.exists(doc.file_path):
        try:
            os.remove(doc.file_path)
        except Exception:
            pass

    if doc.cleaned_file_path and os.path.exists(doc.cleaned_file_path):
        try:
            os.remove(doc.cleaned_file_path)
        except Exception:
            pass

    db.delete(doc)
    db.commit()
    return {"message": "Document deleted successfully"}


# Preload sample helper endpoint (for instant testing or demo)
@app.post("/api/documents/load-sample")
def load_sample_document(
    sample_type: str = Query("pdf", description="pdf, docx, or pptx"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc_id = str(uuid.uuid4())
    st = sample_type.lower()
    
    if st == "pdf":
        filename = "AI_Research_Analysis.pdf"
        file_path = os.path.join(UPLOADS_DIR, f"{doc_id}_{filename}")
        doc = pymupdf.open()
        p = doc.new_page()
        
        # Text with realistic AI indicators: formulaic openers, high regularity, cliches
        sample_body = (
            "Python is a high-level programming language that is widely used for software development. "
            "It provides a simple and readable syntax and supports multiple programming paradigms. "
            "It is important to remember that Python plays a crucial role in modern machine learning systems. "
            "Furthermore, delving into this framework provides a testament to the importance of clean architecture. "
            "In conclusion, modern developers navigate the landscape of automated tools to maximize productivity."
        )
        p.insert_text((50, 72), "COMPREHENSIVE RESEARCH REPORT", fontsize=15)
        p.insert_textbox(pymupdf.Rect(50, 100, 540, 400), sample_body, fontsize=11, fontname="helv")
        
        doc.set_metadata({
            "author": "Dr. Evelyn Vance",
            "title": "Quantum Software Engineering Report",
            "subject": "AI Linguistic Indicators & Metadata Cleaning",
            "creator": "UntraceX PDF Engine 2026",
            "producer": "Research Lab LaTeX Compiler",
            "keywords": "ai, metadata, confidential, document-security",
            "creationDate": "D:20260415103000+05'30'",
            "modDate": "D:20260416184500+05'30'"
        })
        doc.save(file_path)
        doc.close()
        file_type = "PDF"

    elif st == "docx":
        filename = "Corporate_Strategy_Brief.docx"
        file_path = os.path.join(UPLOADS_DIR, f"{doc_id}_{filename}")
        d = docx.Document()
        d.add_heading("Corporate Digital Strategy Briefing", 0)
        d.add_paragraph("Python is a high-level programming language that is widely used for software development.")
        d.add_paragraph("It provides a simple and readable syntax and supports multiple programming paradigms.")
        d.add_paragraph("It is important to remember that our technical infrastructure plays a crucial role in operational excellence.")
        d.add_paragraph("In conclusion, the engineering team must navigate the complexities of digital deployment.")
        d.core_properties.author = "Jonathan Sterling"
        d.core_properties.last_modified_by = "Rebecca Hughes"
        d.core_properties.title = "Q2 Executive Strategy"
        d.core_properties.subject = "Internal Corporate Briefing"
        d.core_properties.keywords = "finance, confidential, strategy, internal-only"
        d.core_properties.comments = "Reviewed by Compliance Team"
        d.core_properties.revision = 7
        d.save(file_path)
        file_type = "DOCX"

    elif st == "pptx":
        filename = "Product_Roadmap_2026.pptx"
        file_path = os.path.join(UPLOADS_DIR, f"{doc_id}_{filename}")
        p = pptx.Presentation()
        slide_layout = p.slide_layouts[0]
        slide = p.slides.add_slide(slide_layout)
        slide.shapes.title.text = "UntraceX Product Strategy"
        slide.placeholders[1].text = "Python is widely used for software development.\nIt provides a simple and readable syntax and supports multiple programming paradigms."
        
        slide2 = p.slides.add_slide(p.slide_layouts[1])
        slide2.shapes.title.text = "Strategic Overview"
        slide2.placeholders[1].text = "It is important to remember that clean documentation plays a crucial role in enterprise security.\nIn conclusion, automated metadata purging ensures privacy."
        
        p.core_properties.author = "Marcus Vance"
        p.core_properties.last_modified_by = "Elena Rostova"
        p.core_properties.title = "UntraceX Product Presentation"
        p.core_properties.subject = "Presentation Slides"
        p.core_properties.keywords = "slides, untracex, presentation, college-demo"
        p.core_properties.revision = 4
        p.save(file_path)
        file_type = "PPTX"

    else:
        filename = "Sample_Notes.txt"
        file_path = os.path.join(UPLOADS_DIR, f"{doc_id}_{filename}")
        content = (
            "Python is a high-level programming language that is widely used for software development. "
            "It provides a simple and readable syntax and supports multiple programming paradigms. "
            "It is important to remember that our technical infrastructure plays a crucial role in operational excellence. "
            "In conclusion, the engineering team must navigate the complexities of digital deployment."
        )
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)
        file_type = "TXT"

    file_size = os.path.getsize(file_path)
    new_doc = Document(
        id=doc_id,
        user_id=current_user.id,
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        file_path=file_path,
        status="Uploaded"
    )
    db.add(new_doc)
    db.commit()
    db.refresh(new_doc)

    return {
        "id": new_doc.id,
        "filename": new_doc.filename,
        "file_type": new_doc.file_type,
        "file_size": new_doc.file_size,
        "status": new_doc.status,
        "created_at": new_doc.created_at.isoformat()
    }
